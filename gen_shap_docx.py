from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

doc = Document()

style = doc.styles['Normal']
font = style.font
font.name = '宋体'
font.size = Pt(11)

# ====== 标题 ======
title = doc.add_heading('SHAP 可解释性分析结果总结', level=0)

# ====== 1. 方法概述 ======
doc.add_heading('1. 方法概述', level=1)
p = doc.add_paragraph()
p.add_run('基于 Stacking 集成模型的 4 个基模型（逻辑回归、随机森林、XGBoost、LightGBM），')
p.add_run('分别计算 SHAP 值，按元模型（逻辑回归）系数的归一化权重进行加权融合，')
p.add_run('得到集成模型的全局特征重要性。')

# ====== 2. 融合权重 ======
doc.add_heading('2. 基模型融合权重', level=1)
table = doc.add_table(rows=5, cols=2, style='Light List Accent 1')
table.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ['基模型', '归一化权重']
for i, h in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.text = h
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True

data = [
    ('XGBoost', '0.2931'),
    ('逻辑回归', '0.2896'),
    ('LightGBM', '0.2790'),
    ('随机森林', '0.1383'),
]
for i, (model, weight) in enumerate(data):
    table.rows[i+1].cells[0].text = model
    table.rows[i+1].cells[1].text = weight

p = doc.add_paragraph()
p.add_run('\n权重分布较为均衡，无单一模型主导决策。XGBoost 贡献最高，随机森林贡献最低，')
p.add_run('与单模型 AUC 性能排序一致。')

# ====== 3. 特征重要性 ======
doc.add_heading('3. 全局特征重要性（Top 10）', level=1)
p = doc.add_paragraph()
p.add_run('已剔除标签泄露特征 chronic_num（慢性病数量）及其衍生列，避免预测偏差。')

table2 = doc.add_table(rows=11, cols=3, style='Light List Accent 1')
table2.alignment = WD_TABLE_ALIGNMENT.CENTER
headers2 = ['排名', '特征', '平均绝对 SHAP 值']
for i, h in enumerate(headers2):
    cell = table2.rows[0].cells[i]
    cell.text = h
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True

data2 = [
    ('1', 'bl_hbalc（糖化血红蛋白）', '0.4115'),
    ('2', 'frailtya（虚弱指数）', '0.1799'),
    ('3', 'bl_glu（空腹血糖）', '0.1075'),
    ('4', 'dyslipe_是（血脂异常）', '0.1069'),
    ('5', 'rgrip（右手握力）', '0.0641'),
    ('6', 'region_东部（东部地区）', '0.0631'),
    ('7', 'frailtyb（虚弱指数b）', '0.0595'),
    ('8', 'bl_ldl（低密度脂蛋白）', '0.0560'),
    ('9', 'hrural_城市_True（城市居住）', '0.0523'),
    ('10', 'retire（退休）', '0.0513'),
]
for i, (rank, feat, val) in enumerate(data2):
    row = table2.rows[i+1]
    row.cells[0].text = rank
    row.cells[1].text = feat
    row.cells[2].text = val

# ====== 4. 分析结论 ======
doc.add_heading('4. 分析结论', level=1)

doc.add_heading('4.1 糖代谢指标主导', level=2)
p = doc.add_paragraph()
p.add_run('bl_hbalc（糖化血红蛋白）和 bl_glu（空腹血糖）分列第 1、3 位，')
p.add_run('合计 SHAP 贡献 0.519，是糖尿病预测最核心的特征，与临床指南一致。')

doc.add_heading('4.2 虚弱与身体功能状态', level=2)
p = doc.add_paragraph()
p.add_run('frailtya（虚弱指数）位列第 2，rgrip（握力）位列第 5，')
p.add_run('说明身体机能状态对中老年糖尿病风险有重要预测价值。')

doc.add_heading('4.3 代谢合并症', level=2)
p = doc.add_paragraph()
p.add_run('dyslipe（血脂异常）、bl_ldl（低密度脂蛋白）进入前十，')
p.add_run('验证了糖脂代谢的病理关联。')

doc.add_heading('4.4 社会人口学因素', level=2)
p = doc.add_paragraph()
p.add_run('region（东部地区）、hrural（城市居住）、retire（退休）等社会因素也具有一定影响力，')
p.add_run('反映生活方式与环境的间接作用。')

save_path = os.path.join(os.path.dirname(__file__), 'SHAP分析结果总结.docx')
doc.save(save_path)
print(f"已保存: {save_path}")
